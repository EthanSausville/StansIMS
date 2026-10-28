"""Export the Stans workbook to JSON.

Usage: python3 tools/excel_export.py Stans.xlsm out.json

Writes two sections:
  inputs   - everything a person types into the workbook (items, costs,
             recipes, production, receipts, counts, sales, settings)
  expected - Excel's own calculated results, used by tests/calc.test.mjs
             to check that the web calculations match the spreadsheet.
"""
import json, sys, datetime, re
import openpyxl

src, out = sys.argv[1], sys.argv[2]
wf = openpyxl.load_workbook(src)                  # formulas
wv = openpyxl.load_workbook(src, data_only=True)  # cached values

def iso(d):
    return d.date().isoformat() if isinstance(d, datetime.datetime) else None

def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None

def val(ws, ref):
    return wv[ws][ref].value

# ---------- items (StandardCost) ----------
sc_f, sc_v = wf['StandardCost'], wv['StandardCost']
years = [sc_v.cell(2, c).value for c in range(5, 15)]
items, std_costs = [], []
pack = {}
for c in range(4, 23):
    item = wf['GoodsReceipt'].cell(2, c).value
    if item: pack[item] = wf['GoodsReceipt'].cell(3, c).value
for r in range(3, sc_f.max_row + 1):
    name = sc_v.cell(r, 1).value
    if not name: continue
    item_no = sc_v.cell(r, 3).value
    typ = sc_v.cell(r, 2).value
    items.append(dict(item_no=item_no, name=name, type=typ, unit=sc_v.cell(r, 4).value,
                      pack_note=pack.get(item_no), sort=r))
    if typ in ('Raw Material', 'Component') or item_no == 10020:
        for i, y in enumerate(years):
            v = num(sc_v.cell(r, 5 + i).value)
            if v is not None:
                std_costs.append(dict(item_no=item_no, year=y, cost=v))
name_to_no = {i['name']: i['item_no'] for i in items}

# ---------- recipes ----------
rc = wv['RecipesCalc']; rcf = wf['RecipesCalc']
recipes, recipe_lines = {}, []
for r in range(2, rc.max_row + 1):
    rname = rc.cell(r, 2).value
    if not rname: continue
    recipe_no = name_to_no['RECIPE: ' + rname]
    recipes.setdefault(recipe_no, dict(item_no=recipe_no, name=rname, batch_size=rc.cell(r, 7).value))
    qf = rcf.cell(r, 5).value
    uses_minutes = isinstance(qf, str) and 'LaborScenarios' in qf
    recipe_lines.append(dict(recipe_item_no=recipe_no, item_no=rc.cell(r, 4).value,
                             qty=None if uses_minutes else rc.cell(r, 5).value,
                             uses_batch_minutes=uses_minutes, sort=r))

# ---------- finished goods bill of materials ----------
fg = wv['FinishedGoodsCalc']
fg_lines = []
for r in range(2, fg.max_row + 1):
    fname = fg.cell(r, 2).value
    if not fname: continue
    fg_lines.append(dict(fg_item_no=name_to_no[fname], item_no=name_to_no[fg.cell(r, 3).value],
                         qty=num(fg.cell(r, 5).value), sort=r))

# ---------- standard prices ----------
sp = wv['StandardPrice']; prices = []
sp_years = [sp.cell(2, c).value for c in range(3, 13)]
for r in range(3, sp.max_row + 1):
    n = sp.cell(r, 1).value
    if not n: continue
    for i, y in enumerate(sp_years):
        v = num(sp.cell(r, 3 + i).value)
        if v is not None: prices.append(dict(item_no=name_to_no[n], year=y, price=v))

# ---------- settings ----------
ls = wv['LaborScenarios']
settings = dict(salaried_count=ls['C2'].value, hourly_a_count=ls['D2'].value, hourly_b_count=ls['E2'].value,
                salaried_wage=ls['G2'].value, hourly_a_rate=ls['H2'].value, hourly_b_rate=ls['I2'].value,
                minutes_per_batch=ls['O1'].value, production_batch_size=wv['Production']['E1'].value)
year_settings = []
for i, y in enumerate(years):
    s = sc_v.cell(1, 5 + i).value
    if isinstance(s, str) and s.startswith('Scenario'):
        year_settings.append(dict(year=y, labor_scenario=int(s.split()[-1])))

# ---------- daily grids ----------
def grid(sheet, item_row, first_row, cols):
    f = wf[sheet]; v = wv[sheet]; rows = []
    for r in range(first_row, f.max_row + 1):
        d = iso(v.cell(r, 3).value)
        if not d: continue
        for c in cols:
            item = f.cell(item_row, c).value
            q = num(v.cell(r, c).value)
            if item and q not in (None, 0) and f.cell(r, c).value is not None:
                rows.append(dict(date=d, item_no=item, qty=q))
    return rows

prod_cols = range(7, 13)
production = []
pf = wf['Production']
fg_by_header = {}
for c in prod_cols:
    h = re.sub(r'\s+', ' ', pf.cell(3, c).value).strip()
    fg_by_header[c] = name_to_no[h]
for r in range(4, pf.max_row + 1):
    d = iso(wv['Production'].cell(r, 3).value)
    for c in prod_cols:
        q = num(wv['Production'].cell(r, c).value)
        if d and q: production.append(dict(date=d, item_no=fg_by_header[c], cases=q))

receipts = grid('GoodsReceipt', 2, 5, list(range(4, 13)) + list(range(14, 23)))
counts = grid('CycleCounts', 2, 5, list(range(4, 13)) + list(range(14, 23)) + list(range(24, 30)))

pdv = wv['PurchaseDemand']; purchase_settings = []
for c in list(range(4, 13)) + list(range(14, 23)):
    purchase_settings.append(dict(item_no=pdv.cell(2, c).value, avg_weekly=pdv.cell(3, c).value,
        max_weekly=pdv.cell(4, c).value, manual_usage=pdv.cell(5, c).value,
        standard_purchase=pdv.cell(7, c).value, lead_time=pdv.cell(8, c).value,
        safety_threshold=pdv.cell(9, c).value))

sr = wv['SalesReport']; sales = []
for r in range(2, sr.max_row + 1):
    d = iso(sr.cell(r, 1).value)
    if not d: continue
    vals = [sr.cell(r, c).value for c in range(2, 8)]
    if all(v is None for v in vals): continue
    inv, q, price, amt, cust, prod = vals
    sales.append(dict(row=r, date=d, invoice=None if inv is None else str(inv), quantity=num(q),
                      price=num(price), amount=num(amt), customer=cust, product=prod,
                      item_no=name_to_no.get(prod)))

# ---------- expected results (Excel cached values) ----------
exp = {}
exp['std_cost_2026'] = {str(sc_v.cell(r, 3).value): sc_v.cell(r, 5).value for r in range(3, 31)}
exp['labor'] = [[ls.cell(r, c).value for c in (11, 12, 13, 15, 16)] for r in range(4, 9)]
pv = wv['Production']
exp['production'] = {iso(pv.cell(r, 3).value): [pv.cell(r, 4).value, pv.cell(r, 5).value, pv.cell(r, 6).value]
                     for r in range(4, pv.max_row + 1) if num(pv.cell(r, 6).value)}
cv = wv['Consumption']; co_items = [cv.cell(2, c).value for c in range(6, 25)]
exp['consumption'] = {}
for r in range(13, cv.max_row + 1):
    d = iso(cv.cell(r, 3).value)
    row = {str(co_items[i]): cv.cell(r, 6 + i).value for i in range(19) if co_items[i] and num(cv.cell(r, 6 + i).value)}
    if d and row: exp['consumption'][d] = row
def daily(sheet, first_row, item_row):
    v = wv[sheet]; f = wf[sheet]; outd = {}
    for r in range(first_row, v.max_row + 1):
        d = iso(v.cell(r, 3).value)
        if not d: continue
        row = {}
        for c in range(4, v.max_column + 1):
            item = f.cell(item_row, c).value
            fv = f.cell(r, c).value
            if item and isinstance(fv, str) and fv.startswith('='):
                row[str(item)] = v.cell(r, c).value
        if row: outd[d] = row
    return outd
exp['daily_inventory'] = daily('DailyInventory', 4, 2)
exp['purchase_demand'] = daily('PurchaseDemand', 12, 2)
ei = wv['EndingInventory$']
exp['ending_inventory'] = {}
for c in range(5, ei.max_column + 1):
    d = ei.cell(5, c).value
    if isinstance(d, datetime.datetime):
        exp['ending_inventory'][iso(d)] = {str(ei.cell(r, 3).value): ei.cell(r, c).value for r in range(8, 32)}
exp['sales'] = {str(r): [sr.cell(r, c).value for c in (10, 11, 12, 13)] for r in range(2, sr.max_row + 1) if iso(sr.cell(r, 1).value)}
dash = wv['YTDDashboard']
exp['dash_sales_by_month'] = [[dash.cell(r, c).value for c in range(27, 31)] for r in range(4, 19)]
exp['dash_loss_by_month'] = [[dash.cell(r, 27).value, dash.cell(r, 28).value] for r in range(42, 57)]

data = dict(inputs=dict(items=items, standard_costs=std_costs, recipes=list(recipes.values()), recipe_lines=recipe_lines,
                        fg_lines=fg_lines, standard_prices=prices, settings=settings, year_settings=year_settings,
                        production=production, goods_receipts=receipts, cycle_counts=counts,
                        purchase_settings=purchase_settings, sales=sales),
            expected=exp)
json.dump(data, open(out, 'w'), indent=1, default=str)
print({k: len(v) if isinstance(v, list) else 1 for k, v in data['inputs'].items()})
